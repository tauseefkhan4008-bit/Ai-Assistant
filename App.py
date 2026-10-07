"""ATS Resume Checker - Streamlit + Google Gemini Flash.

Upload a resume (PDF, DOCX or TXT), optionally paste a job description,
and get an ATS score with concrete suggestions for improvement.
"""

import json
import os
import re
from io import BytesIO

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# ----------------------------------------------------------------------------
# Config
# ----------------------------------------------------------------------------
DEFAULT_MODEL = "gemini-2.5-flash"  # change via the GEMINI_MODEL secret/env var
MAX_RESUME_CHARS = 15000
MAX_JD_CHARS = 6000
MIN_RESUME_CHARS = 100

SECTION_NAMES = [
    "Contact Info",
    "Summary/Objective",
    "Experience",
    "Education",
    "Skills",
    "Keywords",
    "Formatting & Readability",
]

PROMPT_TEMPLATE = """You are an expert ATS (Applicant Tracking System) analyst and professional resume reviewer.

Evaluate the resume below{jd_clause}. Be honest, specific and strict: do not inflate the score.

Scoring guidance:
- 90-100: excellent, ready to submit
- 75-89: good, minor fixes needed
- 55-74: average, clear improvements needed
- below 55: weak, major rework needed
Consider: standard section headings, contact details, quantified achievements, action verbs,
keyword relevance{jd_keywords}, skills coverage, consistent dates, length, and readability
(note that the text was extracted from a file, so ignore odd line breaks).

Return ONLY valid JSON (no markdown, no commentary) with exactly this structure:
{{
  "ats_score": <integer 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "section_scores": {{
    {section_lines}
  }},
  "strengths": ["<short point>", "..."],
  "weaknesses": ["<short point>", "..."],
  "missing_keywords": ["<keyword or skill>", "..."],
  "formatting_issues": ["<short point>", "..."],
  "improvements": [
    {{
      "area": "<section or topic>",
      "issue": "<what is wrong>",
      "suggestion": "<what to do>",
      "example": "<a rewritten example line, or empty string>"
    }}
  ]
}}
Each section score is an integer 0-100. Give 3-6 strengths, 3-6 weaknesses and 5-8 improvements.
Use only information that appears in the resume; never invent experience.

{jd_block}RESUME:
\"\"\"
{resume}
\"\"\"
"""


# ----------------------------------------------------------------------------
# File parsing
# ----------------------------------------------------------------------------
def extract_text(file_name: str, data: bytes) -> str:
    """Extract plain text from PDF, DOCX or TXT bytes."""
    name = file_name.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password protected.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages).strip()
    if name.endswith(".docx"):
        doc = Document(BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip()
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()
    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


# ----------------------------------------------------------------------------
# Prompt + response handling
# ----------------------------------------------------------------------------
def build_prompt(resume_text: str, job_description: str = "") -> str:
    resume_text = resume_text[:MAX_RESUME_CHARS]
    job_description = (job_description or "").strip()[:MAX_JD_CHARS]
    if job_description:
        jd_clause = " against the target job description"
        jd_keywords = " and match with the job description"
        jd_block = f'JOB DESCRIPTION:\n"""\n{job_description}\n"""\n\n'
    else:
        jd_clause = " for general ATS compatibility"
        jd_keywords = ""
        jd_block = ""
    section_lines = ",\n    ".join(f'"{s}": <integer 0-100>' for s in SECTION_NAMES)
    return PROMPT_TEMPLATE.format(
        jd_clause=jd_clause,
        jd_keywords=jd_keywords,
        jd_block=jd_block,
        section_lines=section_lines,
        resume=resume_text,
    )


def parse_json_response(raw: str) -> dict:
    """Parse model output into a dict, tolerating code fences and extra text."""
    if not raw or not raw.strip():
        raise ValueError("The model returned an empty response.")
    text = raw.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*```$", "", text)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        start, end = text.find("{"), text.rfind("}")
        if start != -1 and end > start:
            return json.loads(text[start : end + 1])
        raise ValueError("Could not parse the model response as JSON.")


def _clamp_score(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def _str_list(value) -> list:
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if str(v).strip()]


def normalize_result(data: dict) -> dict:
    """Make sure the result has every field the UI needs, with safe types."""
    if not isinstance(data, dict):
        raise ValueError("Unexpected response format from the model.")

    raw_sections = data.get("section_scores")
    raw_sections = raw_sections if isinstance(raw_sections, dict) else {}
    sections = {name: _clamp_score(raw_sections.get(name)) for name in SECTION_NAMES}

    improvements = []
    for item in data.get("improvements") or []:
        if isinstance(item, dict):
            improvements.append(
                {
                    "area": str(item.get("area", "General")).strip() or "General",
                    "issue": str(item.get("issue", "")).strip(),
                    "suggestion": str(item.get("suggestion", "")).strip(),
                    "example": str(item.get("example", "")).strip(),
                }
            )
        elif isinstance(item, str) and item.strip():
            improvements.append(
                {"area": "General", "issue": "", "suggestion": item.strip(), "example": ""}
            )

    return {
        "ats_score": _clamp_score(data.get("ats_score")),
        "summary": str(data.get("summary", "")).strip(),
        "section_scores": sections,
        "strengths": _str_list(data.get("strengths")),
        "weaknesses": _str_list(data.get("weaknesses")),
        "missing_keywords": _str_list(data.get("missing_keywords")),
        "formatting_issues": _str_list(data.get("formatting_issues")),
        "improvements": improvements,
    }


def score_label(score: int) -> str:
    if score >= 90:
        return "Excellent"
    if score >= 75:
        return "Good"
    if score >= 55:
        return "Average - needs work"
    return "Weak - major rework needed"


# ----------------------------------------------------------------------------
# Gemini
# ----------------------------------------------------------------------------
def get_secret(name: str, default: str = "") -> str:
    """Read from Streamlit secrets first, then environment variables."""
    try:
        if name in st.secrets:
            return str(st.secrets[name])
    except Exception:
        pass  # no secrets file present
    return os.environ.get(name, default)


def analyze_resume(api_key: str, model: str, resume_text: str, job_description: str = "") -> dict:
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=model,
        contents=build_prompt(resume_text, job_description),
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    return normalize_result(parse_json_response(response.text))


# ----------------------------------------------------------------------------
# UI
# ----------------------------------------------------------------------------
def render_result(result: dict) -> None:
    score = result["ats_score"]

    col1, col2 = st.columns([1, 2])
    with col1:
        st.metric("ATS Score", f"{score}/100")
        st.caption(score_label(score))
    with col2:
        st.progress(score / 100)
        if result["summary"]:
            st.write(result["summary"])

    st.subheader("Section scores")
    cols = st.columns(2)
    for i, (name, value) in enumerate(result["section_scores"].items()):
        with cols[i % 2]:
            st.write(f"**{name}** - {value}/100")
            st.progress(value / 100)

    left, right = st.columns(2)
    with left:
        st.subheader("Strengths")
        for s in result["strengths"] or ["No strengths listed."]:
            st.markdown(f"- {s}")
    with right:
        st.subheader("Weaknesses")
        for w in result["weaknesses"] or ["No weaknesses listed."]:
            st.markdown(f"- {w}")

    if result["missing_keywords"]:
        st.subheader("Missing keywords")
        st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))

    if result["formatting_issues"]:
        st.subheader("Formatting issues")
        for f in result["formatting_issues"]:
            st.markdown(f"- {f}")

    st.subheader("How to improve")
    for i, imp in enumerate(result["improvements"], start=1):
        with st.expander(f"{i}. {imp['area']}", expanded=(i <= 3)):
            if imp["issue"]:
                st.markdown(f"**Issue:** {imp['issue']}")
            if imp["suggestion"]:
                st.markdown(f"**Suggestion:** {imp['suggestion']}")
            if imp["example"]:
                st.markdown("**Example:**")
                st.code(imp["example"], language=None)

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="centered")
    st.title("📄 ATS Resume Checker")
    st.write("Upload your resume to get an ATS score and tips to improve it.")

    api_key = get_secret("GEMINI_API_KEY")
    model = get_secret("GEMINI_MODEL", DEFAULT_MODEL)

    with st.sidebar:
        st.header("Settings")
        if not api_key:
            api_key = st.text_input(
                "Gemini API key",
                type="password",
                help="Get a free key at https://aistudio.google.com/apikey",
            )
        else:
            st.success("API key loaded.")
        model = st.text_input("Model", value=model)
        st.caption("Your resume is sent to Google's Gemini API for analysis.")

    uploaded = st.file_uploader("Resume (PDF, DOCX or TXT)", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional)",
        height=160,
        placeholder="Paste the job description for a role-specific score...",
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please add your Gemini API key in the sidebar.")
            return
        try:
            resume_text = extract_text(uploaded.name, uploaded.getvalue())
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            return
        if len(resume_text) < MIN_RESUME_CHARS:
            st.error(
                "Very little text was found. If this is a scanned/image PDF, "
                "export a text-based PDF or upload a DOCX instead."
            )
            return

        with st.spinner("Analyzing your resume..."):
            try:
                result = analyze_resume(api_key, model, resume_text, job_description)
            except Exception as exc:
                st.error(f"Analysis failed: {exc}")
                return

        st.session_state["result"] = result

    if "result" in st.session_state:
        st.divider()
        render_result(st.session_state["result"])


if __name__ == "__main__":
    main()
